# -*- coding: utf-8 -*-
"""验证技能：对 docx/pptx 产物 + 状态文件评分（0-100 分）。

validate(root) -> (score, checks)
  checks: [{id, name, weight, desc, passed, detail}, ...]
"""
import json
import os

from .checklist import (
    CHECKS,
    KEYWORDS,
    PLACEHOLDER_MARKS,
    MIN_PARAGRAPHS,
    MIN_HEADINGS,
    MIN_CHARS,
    MIN_SLIDES,
)


def _find_placeholder(text):
    low = text.lower()
    for m in PLACEHOLDER_MARKS:
        if m.lower() in low:
            return m
    return None


def validate(root):
    results = []

    def add(cid, passed, detail=""):
        c = next(x for x in CHECKS if x["id"] == cid)
        results.append({
            "id": c["id"],
            "name": c["name"],
            "weight": c["weight"],
            "desc": c["desc"],
            "passed": bool(passed),
            "detail": detail,
        })

    docx_path = os.path.join(root, "output", "doc", "report.docx")
    pptx_path = os.path.join(root, "output", "slides", "deck.pptx")
    status_path = os.path.join(root, "status", "loop_status.json")

    # ---------- docx ----------
    docx_ok = False
    doc = None
    if os.path.isfile(docx_path) and os.path.getsize(docx_path) > 0:
        add("docx_exists", True)
        try:
            from docx import Document
            doc = Document(docx_path)
            add("docx_opens", True)
            docx_ok = True
        except Exception as e:  # noqa: BLE001
            add("docx_opens", False, "打开失败: %s" % e)
    else:
        add("docx_exists", False, "文件缺失或为空")

    if docx_ok:
        paras = [p.text for p in doc.paragraphs if p.text.strip()]
        table_texts = [
            cell.text
            for t in doc.tables
            for row in t.rows
            for cell in row.cells
        ]
        all_text = "\n".join(paras + table_texts)
        headings = [
            p for p in doc.paragraphs
            if p.style.name.lower().startswith(("heading", "title"))
        ]
        add("docx_paragraphs", len(paras) >= MIN_PARAGRAPHS,
            "实际 %d / 要求 %d" % (len(paras), MIN_PARAGRAPHS))
        add("docx_headings", len(headings) >= MIN_HEADINGS,
            "实际 %d / 要求 %d" % (len(headings), MIN_HEADINGS))
        add("docx_length", len(all_text) >= MIN_CHARS,
            "实际 %d 字符 / 要求 %d" % (len(all_text), MIN_CHARS))
        missing = [k for k in KEYWORDS if k not in all_text]
        add("docx_keywords", not missing,
            "缺少: " + "/".join(missing) if missing else "关键词齐全")
        cp = doc.core_properties
        add("docx_metadata", bool(cp.title and cp.author),
            "title=%r author=%r" % (cp.title, cp.author))
        mark = _find_placeholder(all_text)
        add("docx_clean", mark is None,
            "发现占位文本 %s" % mark if mark else "无占位文本")
    else:
        for cid in ("docx_paragraphs", "docx_headings", "docx_length",
                    "docx_keywords", "docx_metadata", "docx_clean"):
            add(cid, False, "docx 不可用")

    # ---------- pptx ----------
    pptx_ok = False
    prs = None
    if os.path.isfile(pptx_path) and os.path.getsize(pptx_path) > 0:
        add("pptx_exists", True)
        try:
            from pptx import Presentation
            prs = Presentation(pptx_path)
            add("pptx_opens", True)
            pptx_ok = True
        except Exception as e:  # noqa: BLE001
            add("pptx_opens", False, "打开失败: %s" % e)
    else:
        add("pptx_exists", False, "文件缺失或为空")

    if pptx_ok:
        n = len(prs.slides)
        add("pptx_slides", n >= MIN_SLIDES, "实际 %d / 要求 %d" % (n, MIN_SLIDES))
        first = prs.slides[0] if n else None
        title = first.shapes.title if first is not None else None
        title_ok = title is not None and bool(title.text.strip())
        add("pptx_title", title_ok, "封面标题缺失" if not title_ok else "封面正常")
        texts = [
            shape.text_frame.text
            for slide in prs.slides
            for shape in slide.shapes
            if shape.has_text_frame
        ]
        cp = prs.core_properties
        add("pptx_metadata", bool(cp.title and cp.author),
            "title=%r author=%r" % (cp.title, cp.author))
        mark = _find_placeholder("\n".join(texts))
        add("pptx_clean", mark is None,
            "发现占位文本 %s" % mark if mark else "无占位文本")
    else:
        for cid in ("pptx_slides", "pptx_title", "pptx_metadata", "pptx_clean"):
            add(cid, False, "pptx 不可用")

    # ---------- 状态文件 ----------
    try:
        with open(status_path, encoding="utf-8") as f:
            data = json.load(f)
        ok = "iteration" in data and "score" in data
        add("status_written", ok,
            "iteration=%s score=%s" % (data.get("iteration"), data.get("score")))
    except Exception as e:  # noqa: BLE001
        add("status_written", False, "状态文件不可用: %s" % e)

    score = sum(r["weight"] for r in results if r["passed"])
    return score, results
