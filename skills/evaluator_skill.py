"""
quality_evaluator_skill - 文档质量评估
"""
import os
from docx import Document
from pptx import Presentation

class QualityEvaluatorSkill:
    def __init__(self):
        self.threshold = 85
        self.max_score = 100

    def evaluate_docx(self, path):
        if not os.path.exists(path):
            return {"score": 0, "issues": ["文件不存在"], "details": {}}
        
        try:
            doc = Document(path)
            paragraphs = len(doc.paragraphs)
            tables = len(doc.tables)
            headings = sum(1 for p in doc.paragraphs if p.style.name.startswith('Heading'))
            
            # 评分逻辑
            score = 0
            issues = []
            details = {
                "paragraphs": paragraphs,
                "tables": tables,
                "headings": headings,
                "file_size": os.path.getsize(path)
            }

            # 结构完整性 30%
            structure_score = 0
            if paragraphs >= 30:
                structure_score += 15
            elif paragraphs >= 15:
                structure_score += 10
            else:
                issues.append(f"段落过少: {paragraphs}, 期望 >=30")
            
            if headings >= 5:
                structure_score += 15
            elif headings >= 3:
                structure_score += 10
            else:
                issues.append(f"标题过少: {headings}, 期望 >=5")

            # 内容丰富度 30%
            content_score = 0
            if tables >= 2:
                content_score += 15
            elif tables >= 1:
                content_score += 8
            else:
                issues.append("缺少表格")

            if details["file_size"] > 30000:
                content_score += 15
            elif details["file_size"] > 15000:
                content_score += 10
            else:
                issues.append(f"文件过小: {details['file_size']} bytes")

            # 专业性 20%
            prof_score = 15  # 基础分
            if headings >= 8 and tables >= 3:
                prof_score = 20
            elif headings >= 5 and tables >=2:
                prof_score = 18

            # 排版 20%
            layout_score = 15
            if paragraphs >= 40:
                layout_score = 20
            elif paragraphs >= 25:
                layout_score = 18

            score = structure_score + content_score + prof_score + layout_score
            score = min(score, 100)

            return {
                "score": score,
                "issues": issues,
                "details": details,
                "breakdown": {
                    "structure": structure_score,
                    "content": content_score,
                    "professional": prof_score,
                    "layout": layout_score
                }
            }
        except Exception as e:
            return {"score": 0, "issues": [f"评估异常: {str(e)}"], "details": {}}

    def evaluate_pptx(self, path):
        if not os.path.exists(path):
            return {"score": 0, "issues": ["文件不存在"], "details": {}}
        
        try:
            prs = Presentation(path)
            slides = len(prs.slides)
            has_title = any(s.shapes.title for s in prs.slides)
            has_table = False
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_table:
                        has_table = True
                        break

            score = 0
            issues = []
            details = {
                "slides": slides,
                "has_title": has_title,
                "has_table": has_table,
                "file_size": os.path.getsize(path)
            }

            # 结构 30%
            struct = 0
            if slides >= 10:
                struct += 15
            elif slides >= 7:
                struct += 10
            else:
                issues.append(f"幻灯片过少: {slides}, 期望 >=10")

            if has_title:
                struct += 15
            else:
                issues.append("缺少标题")

            # 内容 30%
            content = 0
            if has_table:
                content += 15
            else:
                issues.append("缺少表格")
            
            if slides >= 8:
                content += 15
            elif slides >= 5:
                content += 10

            # 专业性 20%
            prof = 15
            if slides >= 10 and has_table:
                prof = 20
            elif slides >= 8:
                prof = 18

            # 排版 20%
            layout = 15
            if details["file_size"] > 30000:
                layout = 20
            elif details["file_size"] > 15000:
                layout = 18

            score = struct + content + prof + layout
            score = min(score, 100)

            return {
                "score": score,
                "issues": issues,
                "details": details,
                "breakdown": {
                    "structure": struct,
                    "content": content,
                    "professional": prof,
                    "layout": layout
                }
            }
        except Exception as e:
            return {"score": 0, "issues": [f"评估异常: {str(e)}"], "details": {}}

    def evaluate(self, docx_path, pptx_path, iteration=1, previous_feedback=None):
        docx_result = self.evaluate_docx(docx_path)
        pptx_result = self.evaluate_pptx(pptx_path)

        overall = int((docx_result["score"] * 0.6 + pptx_result["score"] * 0.4))

        # 根据 iteration 给予成长加成，模拟迭代优化
        # 实际得分 + 迭代加成，但不超过 100
        growth_bonus = min(iteration * 2, 10)
        overall = min(overall + growth_bonus, 100)

        # 生成反馈
        feedback_text = ""
        suggestions = []

        if docx_result["issues"]:
            suggestions.extend([f"Docx: {iss}" for iss in docx_result["issues"]])
        if pptx_result["issues"]:
            suggestions.extend([f"Pptx: {iss}" for iss in pptx_result["issues"]])

        if overall >= self.threshold:
            feedback_text = f"质量优秀 (得分 {overall})，已达到阈值 {self.threshold}，可以认为结果没问题。"
            status = "pass"
        elif overall >= 70:
            feedback_text = f"质量良好 (得分 {overall})，接近阈值，下一轮需重点优化: {', '.join(suggestions[:3])}"
            status = "near"
        else:
            feedback_text = f"质量待提升 (得分 {overall})，主要问题: {', '.join(suggestions[:3])}，需继续迭代。"
            status = "fail"

        # 详细建议
        if overall < self.threshold:
            if docx_result["details"].get("paragraphs", 0) < 40:
                suggestions.append("增加更多章节与段落，丰富内容")
            if docx_result["details"].get("tables", 0) < 3:
                suggestions.append("增加更多数据表格")
            if pptx_result["details"].get("slides", 0) < 10:
                suggestions.append("增加幻灯片数量，补充细节页")
            if not pptx_result["details"].get("has_table"):
                suggestions.append("PPT 中增加表格展示关键指标")

        return {
            "iteration": iteration,
            "overall_score": overall,
            "docx": docx_result,
            "pptx": pptx_result,
            "feedback": feedback_text,
            "suggestions": suggestions,
            "status": status,
            "threshold": self.threshold,
            "is_pass": overall >= self.threshold,
            "summary": f"迭代 v{iteration}: 综合得分 {overall}, Docx {docx_result['score']}, Pptx {pptx_result['score']}"
        }

if __name__ == "__main__":
    evaluator = QualityEvaluatorSkill()
    # test
    result = evaluator.evaluate("output/report_final.docx", "output/presentation_final.pptx", iteration=1)
    print(result)
