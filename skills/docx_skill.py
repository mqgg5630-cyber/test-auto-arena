# -*- coding: utf-8 -*-
"""docx 生成技能：按内容模型 + 上轮验证失败项生成 Word 报告。

build_docx(path, level=1, fixes=(), history=(), result=None)
  level   : 内容完整度 1=初稿 / 2=修订 / 3=完整
  fixes   : 上一轮验证失败项集合（驱动自修复）
  history : 迭代历史（渲染进「迭代记录」表）
  result  : 本轮验证结果 dict(score, checks)；产出最终版时传入并写入文档
"""
import os

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from .checklist import CHECKS
from pipeline import content_spec as spec


def _set_base_font(doc):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "微软雅黑")


def _add_section(doc, sec):
    doc.add_heading(sec["heading"], level=1)
    for para in sec["body"]:
        doc.add_paragraph(para)


def _add_validation_table(doc, result):
    doc.add_paragraph("流水线对每轮产物执行以下 15 项检查，得分 = 通过项权重之和（满分 100）：")
    table = doc.add_table(rows=len(CHECKS) + 1, cols=4)
    table.style = "Table Grid"
    for i, h in enumerate(["检查项", "权重", "说明", "本轮结果"]):
        table.rows[0].cells[i].text = h
    res_map = {c["id"]: c for c in (result["checks"] if result else [])}
    for i, c in enumerate(CHECKS, start=1):
        cells = table.rows[i].cells
        cells[0].text = c["name"]
        cells[1].text = str(c["weight"])
        cells[2].text = c["desc"]
        if result:
            r = res_map.get(c["id"])
            cells[3].text = "PASS" if (r and r["passed"]) else "FAIL"
        else:
            cells[3].text = "待本轮验证"


def _add_iterations_table(doc, history):
    doc.add_paragraph("各轮迭代的得分与状态（由 pipeline/loop.py 自动记录并推送）：")
    if not history:
        doc.add_paragraph("（首版文档，尚无历史迭代）")
        return
    table = doc.add_table(rows=len(history) + 1, cols=5)
    table.style = "Table Grid"
    for i, h in enumerate(["迭代", "得分", "失败项", "状态", "推送"]):
        table.rows[0].cells[i].text = h
    for i, h_ in enumerate(history, start=1):
        cells = table.rows[i].cells
        cells[0].text = str(h_["iteration"])
        cells[1].text = "%s/100" % h_["score"]
        cells[2].text = str(len(h_["failed"]))
        cells[3].text = h_["status"]
        cells[4].text = h_.get("push", "")


def _add_conclusion(doc, result, history):
    if result:
        checks = result["checks"]
        passed = sum(1 for c in checks if c["passed"])
        doc.add_paragraph("本轮验证得分 %s/100，%d/%d 项检查通过。"
                          % (result["score"], passed, len(checks)))
        if result["score"] >= 100:
            doc.add_paragraph("全部验证项通过：流水线判定结果收敛，自动循环正常终止。")
        else:
            bad = [c["name"] for c in checks if not c["passed"]]
            doc.add_paragraph("仍未通过：" + "、".join(bad) + "，下一轮将针对性修复。")
    doc.add_paragraph("产物位置：output/doc/report.docx（Word 报告）、output/slides/deck.pptx（PPT 演示）。")
    doc.add_paragraph("状态追溯：status/loop_status.json（机器可读）、status/REPORT.md（人类可读），每轮迭代均提交并推送到分支。")
    if result is None:
        doc.add_paragraph("本版本为自循环自动修订版：依据上一轮验证失败项针对性修复生成，本轮最终验证结果以 status/ 目录与分支提交记录为准。")


def build_docx(path, level=1, fixes=(), history=(), result=None):
    fixes = set(fixes)
    full = level >= 3 or bool(fixes & spec.DOCX_CONTENT_FIXES)
    meta = level >= 2 or "docx_metadata" in fixes
    clean = level >= 2 or "docx_clean" in fixes

    doc = Document()
    _set_base_font(doc)
    doc.add_heading(spec.TITLE, level=0)
    doc.add_paragraph(spec.SUBTITLE)

    for sec in spec.SECTIONS[:2]:
        _add_section(doc, sec)

    if not full:
        # 初稿：仅到「流程说明」为止，其余章节与元数据待自循环补齐
        doc.add_heading(spec.SECTIONS[2]["heading"], level=1)
        doc.add_paragraph("（初稿：流程说明将在后续迭代中由自循环补齐）")
        doc.add_paragraph(spec.PLACEHOLDER if not clean else spec.PLACEHOLDER_CLEAN)
    else:
        for sec in spec.SECTIONS[2:]:
            _add_section(doc, sec)
            if sec["key"] == "validation":
                _add_validation_table(doc, result)
            elif sec["key"] == "iterations":
                _add_iterations_table(doc, history)
            elif sec["key"] == "conclusion":
                _add_conclusion(doc, result, history)

    if meta:
        cp = doc.core_properties
        cp.title = spec.TITLE
        cp.author = spec.AUTHOR
        cp.subject = spec.SUBTITLE
        cp.comments = "auto-loop pipeline 自动生成" + ("（最终版）" if result else "")

    os.makedirs(os.path.dirname(path), exist_ok=True)
    doc.save(path)
    return path
